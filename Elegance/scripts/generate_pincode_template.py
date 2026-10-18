"""Builds pincode_rates_template.xlsx from a pincode -> state dataset.

The dataset is a CSV with the columns `key` (e.g. IN/110001) and `admin_name1`
(state), such as the public India postal-code dump. Consecutive pincodes that
belong to the same state are collapsed into one range, so the generated ranges
never overlap.

Usage:
    python scripts/generate_pincode_template.py --source IN.csv
    python scripts/generate_pincode_template.py --source IN.csv --output my_rates.xlsx

Prices in the template are sensible zone defaults. Edit them in Excel before
running import_pincode_rates.py.
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DEFAULT_PRICE = 99

# Dataset spellings -> the names customers know.
STATE_ALIASES = {
    "Andaman and Nicobar": "Andaman and Nicobar Islands",
    "Chattisgarh": "Chhattisgarh",
    "Dadar and Nagar Haveli": "Dadra and Nagar Haveli",
    "Daman and Diu": "Daman and Diu",
    "New Delhi": "Delhi",
    "Orisa": "Odisha",
    "Pondicherry": "Puducherry",
    "Uttaranchal": "Uttarakhand",
}

ZONE_PRICES = {
    # West
    "Gujarat": 59, "Maharashtra": 69, "Goa": 79, "Daman and Diu": 79, "Dadra and Nagar Haveli": 79,
    "Rajasthan": 79, "Madhya Pradesh": 89,
    # North
    "Delhi": 89, "Haryana": 89, "Punjab": 99, "Chandigarh": 99, "Uttar Pradesh": 99,
    "Uttarakhand": 109, "Himachal Pradesh": 119,
    # South
    "Karnataka": 89, "Telangana": 99, "Andhra Pradesh": 99, "Tamil Nadu": 99, "Kerala": 109,
    "Puducherry": 109,
    # East / Central
    "Chhattisgarh": 109, "Odisha": 119, "Jharkhand": 119, "Bihar": 119, "West Bengal": 119,
    # North-East, hills and islands
    "Assam": 149, "Sikkim": 149, "Meghalaya": 159, "Tripura": 159, "Manipur": 159, "Mizoram": 159,
    "Nagaland": 159, "Arunachal Pradesh": 169, "Jammu and Kashmir": 149, "Ladakh": 169,
    "Andaman and Nicobar Islands": 249, "Lakshadweep": 249,
}


def load_ranges(source):
    pincodes = []
    with open(source, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            digits = row["key"].split("/")[-1].strip()
            if not (digits.isdigit() and len(digits) == 6):
                continue
            state = STATE_ALIASES.get(row["admin_name1"].strip(), row["admin_name1"].strip())
            pincodes.append((int(digits), state))

    pincodes = sorted(set(pincodes))

    # A pincode listed under two states would make overlapping ranges, so keep the first.
    seen, unique = set(), []
    for number, state in pincodes:
        if number not in seen:
            seen.add(number)
            unique.append((number, state))

    ranges = []
    for number, state in unique:
        if ranges and ranges[-1][2] == state:
            ranges[-1][1] = number
        else:
            ranges.append([number, number, state])
    return ranges


def build_workbook(ranges):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pincode Rates"

    headers = ["State", "Pincode Range", "Price", "Active"]
    sheet.append(headers)
    for start, end, state in ranges:
        sheet.append([state, f"{start}-{end}", ZONE_PRICES.get(state, DEFAULT_PRICE), "Yes"])

    header_fill = PatternFill("solid", fgColor="1F1F1F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    for column, width in zip("ABCD", (30, 20, 12, 10)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "@"  # keep ranges as text so Excel never reformats them

    active_validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    sheet.add_data_validation(active_validation)
    active_validation.add(f"D2:D{sheet.max_row}")

    guide = workbook.create_sheet("How to use")
    for line in (
        "Columns",
        "State          Label only (used for display and filtering). Price is decided by the pincode range.",
        "Pincode Range  6-digit start and end, like 380001-396590. A single pincode (110001) is also fine.",
        "Price          Delivery charge in rupees (0 = free delivery for that range).",
        "Active         Optional. Yes/No. Inactive ranges are kept but not deliverable. Defaults to Yes.",
        "",
        "Rules",
        "Ranges must not overlap each other, in this file or with ranges already in the database.",
        "Rows with the exact same range as an existing one UPDATE that range (state, price, active).",
        "To resize or remove ranges using this file, run the importer with --prune (see --help).",
        "",
        "Run:  python scripts/import_pincode_rates.py pincode_rates.xlsx --dry-run",
        "Then: python scripts/import_pincode_rates.py pincode_rates.xlsx",
    ):
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    for cell in (guide["A1"], guide["A7"]):
        cell.font = Font(bold=True)
    return workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source", required=True, help="CSV with `key` and `admin_name1` columns")
    parser.add_argument("--output", default=str(Path(__file__).with_name("pincode_rates_template.xlsx")))
    args = parser.parse_args()

    ranges = load_ranges(args.source)
    build_workbook(ranges).save(args.output)
    states = len({state for _, _, state in ranges})
    print(f"Wrote {len(ranges)} ranges across {states} states to {args.output}")


if __name__ == "__main__":
    main()
