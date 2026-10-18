"""Imports delivery rates from an Excel file into the database.

Expected columns (header names are case-insensitive, extra columns are ignored):
    State | Pincode Range | Price | Active (optional)

Behaviour
  * Every row goes through the same validation as the admin page and API
    (6-digit pincodes, start <= end, no overlap with anything else).
  * A row whose range exactly matches an existing one updates it in place,
    so the script is safe to re-run whenever prices change.
  * Nothing is written if any row is invalid, unless --skip-invalid is given.
  * Ranges that exist in the database but not in the file are left alone.
    Use --prune to make the file the single source of truth (this is also how
    to resize a range: change it in the file and run with --prune).

Usage:
    python scripts/import_pincode_rates.py pincode_rates.xlsx --dry-run
    python scripts/import_pincode_rates.py pincode_rates.xlsx
    python scripts/import_pincode_rates.py pincode_rates.xlsx --prune
"""

import argparse
import os
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "Elegance.settings")

import django  # noqa: E402

django.setup()

from django.core.exceptions import ValidationError  # noqa: E402
from django.db import transaction  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from Order.models import PincodeRange  # noqa: E402
from Order.shipping import format_range, normalize_state, parse_pincode_range  # noqa: E402

COLUMN_ALIASES = {
    "state": "state",
    "pincode range": "range",
    "pincode_range": "range",
    "pincode": "range",
    "range": "range",
    "price": "price",
    "delivery price": "price",
    "active": "active",
    "is_active": "active",
}


class Rollback(Exception):
    pass


def read_rows(path):
    sheet = load_workbook(path, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)

    header = next(rows, None) or ()
    columns = {}
    for index, name in enumerate(header):
        key = COLUMN_ALIASES.get(str(name or "").strip().lower())
        if key and key not in columns:
            columns[key] = index

    missing = [label for label, key in (("State", "state"), ("Pincode Range", "range"), ("Price", "price")) if key not in columns]
    if missing:
        sys.exit(f"Missing column(s): {', '.join(missing)}. Expected headers: State, Pincode Range, Price.")

    for number, row in enumerate(rows, start=2):
        values = {key: (row[index] if index < len(row) else None) for key, index in columns.items()}
        if all(value in (None, "") for value in values.values()):
            continue
        yield number, values


def parse_price(value):
    try:
        price = Decimal(str(value).strip().replace(",", "")).quantize(Decimal("0.01"))
    except (InvalidOperation, AttributeError):
        raise ValueError(f"'{value}' is not a valid price.")
    if price < 0:
        raise ValueError("Price cannot be negative.")
    if price >= Decimal("1000000"):
        raise ValueError("Price is too large.")
    return price


def parse_active(value):
    if value in (None, ""):
        return True
    text = str(value).strip().lower()
    if text in ("yes", "y", "true", "1", "active"):
        return True
    if text in ("no", "n", "false", "0", "inactive"):
        return False
    raise ValueError(f"Active must be Yes or No, got '{value}'.")


def parse_row(values):
    state = normalize_state(values.get("state"))
    if not state:
        raise ValueError("State is required.")
    start, end = parse_pincode_range(values.get("range"))
    return state, start, end, parse_price(values.get("price")), parse_active(values.get("active"))


def run(path, dry_run, prune, skip_invalid):
    parsed, errors = [], []
    for number, values in read_rows(path):
        try:
            parsed.append((number, *parse_row(values)))
        except ValueError as error:
            errors.append((number, str(error)))

    created = updated = unchanged = deleted = 0

    try:
        with transaction.atomic():
            if prune:
                keep = {(start, end) for _, _, start, end, _, _ in parsed}
                for existing in PincodeRange.objects.all():
                    if (existing.pincode_start, existing.pincode_end) not in keep:
                        existing.delete()
                        deleted += 1

            for number, state, start, end, price, active in parsed:
                instance = PincodeRange.objects.filter(pincode_start=start, pincode_end=end).first()
                is_new = instance is None
                if is_new:
                    instance = PincodeRange(pincode_start=start, pincode_end=end)
                elif (instance.state, instance.price, instance.is_active) == (state, price, active):
                    unchanged += 1
                    continue

                instance.state, instance.price, instance.is_active = state, price, active
                try:
                    instance.full_clean()
                    instance.save()
                except ValidationError as error:
                    errors.append((number, " ".join(error.messages)))
                    continue

                if is_new:
                    created += 1
                else:
                    updated += 1

            if (errors and not skip_invalid) or dry_run:
                raise Rollback()
    except Rollback:
        pass

    return created, updated, unchanged, deleted, errors


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("file", help="Path to the .xlsx file")
    parser.add_argument("--dry-run", action="store_true", help="Validate and report, but write nothing")
    parser.add_argument("--prune", action="store_true", help="Delete database ranges that are not in the file")
    parser.add_argument("--skip-invalid", action="store_true", help="Import valid rows even if some rows fail")
    args = parser.parse_args()

    if not Path(args.file).is_file():
        sys.exit(f"File not found: {args.file}")

    created, updated, unchanged, deleted, errors = run(args.file, args.dry_run, args.prune, args.skip_invalid)

    for number, message in errors:
        print(f"  Row {number}: {message}")

    blocked = bool(errors) and not args.skip_invalid
    prefix = "[dry run] " if args.dry_run else ("[nothing written] " if blocked else "")
    print(
        f"{prefix}created {created}, updated {updated}, unchanged {unchanged}"
        + (f", deleted {deleted}" if args.prune else "")
        + f", errors {len(errors)}"
    )
    if blocked:
        print("Fix the rows above and run again, or pass --skip-invalid to import the valid rows only.")
        sys.exit(1)


if __name__ == "__main__":
    main()
